import io
import re
import zipfile
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document

from core.config import Settings, get_settings
from core.models import ResumeDocument
from core.redaction import redact_pii
from core.security import sanitize_extracted_text


class ResumeIngestionError(ValueError):
    pass


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
_PDF_ACTIVE_CONTENT = re.compile(
    rb"/(?:JavaScript|JS|Launch|EmbeddedFile|RichMedia)\b", re.IGNORECASE
)
_DOCX_ACTIVE_PATH_PARTS = (
    "vbaproject.bin",
    "/embeddings/",
    "/activex/",
    "/externallinks/",
    "/oleobject",
    "customui/",
)


def _near_white(color: Any) -> bool:
    if color is None:
        return False
    values = list(color) if isinstance(color, (list, tuple)) else [color]
    try:
        numeric = [float(value) for value in values[:3]]
    except (TypeError, ValueError):
        return False
    if not numeric:
        return False
    threshold = 0.95 if max(numeric) <= 1.0 else 245.0
    return all(value >= threshold for value in numeric)


def _visible_pdf_char(character: dict[str, Any], width: float, height: float) -> bool:
    try:
        if float(character.get("size", 12)) < 3:
            return False
        x0 = float(character.get("x0", 0))
        x1 = float(character.get("x1", x0))
        top = float(character.get("top", 0))
        bottom = float(character.get("bottom", top))
        if x1 < 0 or x0 > width or bottom < 0 or top > height:
            return False
    except (TypeError, ValueError):
        return False
    return not _near_white(character.get("non_stroking_color"))


def _parse_pdf(content: bytes, settings: Settings) -> tuple[str, list[str]]:
    if b"%PDF-" not in content[:1024]:
        raise ResumeIngestionError("The file extension is PDF, but its signature is not valid.")
    if _PDF_ACTIVE_CONTENT.search(content):
        raise ResumeIngestionError(
            "PDFs containing scripts, embedded files, or active content are not allowed."
        )

    flags: list[str] = []
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            if len(pdf.pages) > settings.max_pdf_pages:
                raise ResumeIngestionError(
                    f"PDF exceeds the {settings.max_pdf_pages}-page safety limit."
                )
            page_text: list[str] = []
            removed_hidden = False
            for page in pdf.pages:
                page_width = page.width
                page_height = page.height
                visible_chars = [
                    character
                    for character in page.chars
                    if _visible_pdf_char(character, page_width, page_height)
                ]
                if len(visible_chars) != len(page.chars):
                    removed_hidden = True
                filtered = page.filter(
                    lambda obj, width=page_width, height=page_height: (
                        obj.get("object_type") != "char" or _visible_pdf_char(obj, width, height)
                    )
                )
                page_text.append(filtered.extract_text() or "")
            if removed_hidden:
                flags.append("hidden_or_invisible_pdf_text_removed")
            text = "\n".join(page_text)
    except ResumeIngestionError:
        raise
    except Exception as exc:
        raise ResumeIngestionError(f"Could not read PDF: {exc}") from exc
    if not text.strip():
        raise ResumeIngestionError(
            "This PDF has no visible extractable text and likely needs OCR before it can be scored."
        )
    return text, flags


def _inspect_docx_package(content: bytes, settings: Settings) -> None:
    if not content.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
        raise ResumeIngestionError(
            "The file extension is DOCX, but its ZIP signature is not valid."
        )
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > settings.max_docx_members:
                raise ResumeIngestionError("DOCX contains too many package entries.")
            names = {member.filename for member in members}
            if "word/document.xml" not in names or "[Content_Types].xml" not in names:
                raise ResumeIngestionError("DOCX package is missing required document parts.")

            total_uncompressed = 0
            for member in members:
                normalized_name = member.filename.replace("\\", "/")
                lowered_name = f"/{normalized_name.lower()}"
                path = Path(normalized_name)
                if path.is_absolute() or ".." in path.parts:
                    raise ResumeIngestionError("DOCX contains an unsafe package path.")
                if member.flag_bits & 0x1:
                    raise ResumeIngestionError("Encrypted DOCX package entries are not allowed.")
                if any(part in lowered_name for part in _DOCX_ACTIVE_PATH_PARTS):
                    raise ResumeIngestionError(
                        "DOCX files containing macros, embedded objects, or active content are "
                        "not allowed."
                    )
                total_uncompressed += member.file_size
                if total_uncompressed > settings.max_docx_uncompressed_bytes:
                    raise ResumeIngestionError(
                        "DOCX expands beyond the safe uncompressed-size limit."
                    )
                if member.file_size and (
                    member.file_size / max(member.compress_size, 1)
                    > settings.max_docx_compression_ratio
                ):
                    raise ResumeIngestionError("DOCX has an unsafe compression ratio.")
                if normalized_name.lower().endswith((".xml", ".rels")):
                    markup = archive.read(member)
                    lowered_markup = markup.lower()
                    if b"<!doctype" in lowered_markup or b"<!entity" in lowered_markup:
                        raise ResumeIngestionError(
                            "DOCX XML entities and document types are not allowed."
                        )
    except ResumeIngestionError:
        raise
    except (zipfile.BadZipFile, OSError, RuntimeError) as exc:
        raise ResumeIngestionError(f"Could not inspect DOCX package: {exc}") from exc


def _font_is_hidden(font: Any) -> bool:
    if font.hidden is True:
        return True
    if font.size is not None and font.size.pt < 3:
        return True
    rgb = font.color.rgb
    return rgb is not None and str(rgb).upper() == "FFFFFF"


def _run_is_hidden(run: Any) -> bool:
    fonts = [run.font]
    if run.style is not None:
        fonts.append(run.style.font)
    paragraph = getattr(run, "_parent", None)
    if paragraph is not None and paragraph.style is not None:
        fonts.append(paragraph.style.font)
    return any(_font_is_hidden(font) for font in fonts)


def _paragraph_visible_text(paragraph: Any) -> tuple[str, bool]:
    if not paragraph.runs:
        return paragraph.text, False
    visible: list[str] = []
    removed = False
    for run in paragraph.runs:
        if _run_is_hidden(run):
            removed = removed or bool(run.text)
        else:
            visible.append(run.text)
    return "".join(visible), removed


def _parse_docx(content: bytes, settings: Settings) -> tuple[str, list[str]]:
    _inspect_docx_package(content, settings)
    try:
        document = Document(io.BytesIO(content))
        lines: list[str] = []
        removed_hidden = False
        for paragraph in document.paragraphs:
            text, removed = _paragraph_visible_text(paragraph)
            lines.append(text)
            removed_hidden = removed_hidden or removed
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        text, removed = _paragraph_visible_text(paragraph)
                        lines.append(text)
                        removed_hidden = removed_hidden or removed
        flags = ["hidden_or_white_docx_text_removed"] if removed_hidden else []
        return "\n".join(lines), flags
    except Exception as exc:
        raise ResumeIngestionError(f"Could not read DOCX: {exc}") from exc


def _validate_and_normalize_text(
    text: str, settings: Settings, flags: list[str] | None = None
) -> tuple[str, list[str]]:
    if "\x00" in text:
        raise ResumeIngestionError("Resume text contains binary NUL bytes.")
    normalized, normalization_flags = sanitize_extracted_text(text)
    combined_flags = list(dict.fromkeys([*(flags or []), *normalization_flags]))
    normalized = normalized.strip()
    if not normalized:
        raise ResumeIngestionError("Resume contains no readable text.")
    if len(normalized) > settings.max_resume_text_chars:
        raise ResumeIngestionError(
            f"Resume text exceeds the {settings.max_resume_text_chars:,}-character limit."
        )
    return normalized, combined_flags


def parse_resume_text(text: str, name: str, settings: Settings | None = None) -> ResumeDocument:
    settings = settings or get_settings()
    normalized, flags = _validate_and_normalize_text(text, settings)
    return ResumeDocument(
        name=Path(name).name,
        text=normalized,
        redacted_text=redact_pii(normalized),
        size_bytes=len(text.encode("utf-8")),
        security_flags=flags,
    )


def parse_resume_bytes(
    content: bytes, filename: str, settings: Settings | None = None
) -> ResumeDocument:
    settings = settings or get_settings()
    if len(content) > settings.max_file_bytes:
        limit_mb = settings.max_file_bytes / 1024 / 1024
        raise ResumeIngestionError(f"Resume exceeds the {limit_mb:g} MB file-size limit.")
    if not content:
        raise ResumeIngestionError("Resume file is empty.")

    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ResumeIngestionError("Unsupported resume type. Use PDF, DOCX, or TXT.")
    if suffix == ".pdf":
        text, flags = _parse_pdf(content, settings)
    elif suffix == ".docx":
        text, flags = _parse_docx(content, settings)
    else:
        if b"\x00" in content:
            raise ResumeIngestionError("TXT resume appears to contain binary data.")
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ResumeIngestionError("TXT resumes must use UTF-8 encoding.") from exc
        flags = []

    normalized, flags = _validate_and_normalize_text(text, settings, flags)
    return ResumeDocument(
        name=Path(filename).name,
        text=normalized,
        redacted_text=redact_pii(normalized),
        size_bytes=len(content),
        security_flags=flags,
    )


def parse_resume_file(path: Path | str, settings: Settings | None = None) -> ResumeDocument:
    path = Path(path)
    if not path.is_file():
        raise ResumeIngestionError(f"Resume not found: {path}")
    document = parse_resume_bytes(path.read_bytes(), path.name, settings)
    document.source_path = str(path)
    return document
